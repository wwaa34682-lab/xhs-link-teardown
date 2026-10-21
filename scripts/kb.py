#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""对标库：建库 / 检查 / 只追加写入。
  python kb.py init  [--backend csv|xlsx|feishu] [--name 对标拆解库]   # 没有库时按标准结构新建
  python kb.py check                                                 # 读出目标库的列，显示映射（只读）
  python kb.py write <帖子文件夹> [...] [--dry-run]                    # 拆解完成后追加写入（按记录ID去重）
配置：config/kb_local.yaml（不存在则用 kb_local.example.yaml 的默认 = 本地 CSV）。
硬规则：只新增行；不改、不删已有行；不改已有表的字段/结构；记录ID 已存在就跳过并报告。
"""
import argparse, csv, json, os, re, sys, time, urllib.request, importlib.util
from pathlib import Path
for s in (sys.stdout, sys.stderr):
    try: s.reconfigure(encoding="utf-8", errors="replace")
    except Exception: pass
import yaml  # pip install pyyaml

ROOT = Path(__file__).resolve().parent.parent
PENDING = "【待拆解】"

def load_cfg(p=None):
    p = Path(p) if p else ROOT / "config" / "kb_local.yaml"
    if not p.exists(): p = ROOT / "config" / "kb_local.example.yaml"
    c = yaml.safe_load(open(p, encoding="utf-8")) or {}
    c["_path"] = str(p); return c

def schema():
    return yaml.safe_load(open(ROOT / "config" / "kb_schema.yaml", encoding="utf-8"))

def all_fields(S):
    return [f["name"] for f in S["core_fields"]] + list(S["teardown_fields"])

# ---------- 从帖子文件夹组装一行 ----------
def parse_a2(md):
    d = {}
    for m in re.finditer(r"^\|\s*\d+\s*\|\s*([^|]+?)\s*\|\s*(.*?)\s*\|\s*$", md, re.M):
        d[m.group(1)] = m.group(2)
    return d

def section(md, title):
    m = re.search(rf"^## {re.escape(title)}\s*\n(.*?)(?=^## |^---|\Z)", md, re.M | re.S)
    return m.group(1).strip() if m else ""

def build_row(folder, cfg, S):
    folder = Path(folder)
    md = (folder / "拆解.md").read_text(encoding="utf-8")
    if PENDING in md:
        n = md.count(PENDING)
        raise ValueError(f"{folder.name}: 拆解.md 还有 {n} 处 {PENDING}，先补全再入库")
    meta = json.loads((folder / "meta.json").read_text(encoding="utf-8")) if (folder / "meta.json").exists() else {}
    a2 = parse_a2(md)
    brand = cfg.get("brand", "本品牌")
    nid = meta.get("作品ID") or a2.get("样本ID", "").replace("BM-", "") or "待补"
    imgs = sorted(p.name for p in folder.glob("[0-9][0-9].jpg"))
    urls = [u for u in (meta.get("下载地址") or []) if u]
    day = str(meta.get("抓取日期", time.strftime("%Y%m%d")))
    day = f"{day[:4]}-{day[4:6]}-{day[6:8]}" if re.fullmatch(r"\d{8}", day) else day
    def g(k, *alts):
        for x in (k,) + alts:
            if a2.get(x): return a2[x]
        return "待补"
    likes = re.search(r"互动数据（采集时）：([^。]+)", md)
    row = {
        S["id_field"]: S.get("id_prefix", "BM-XHS-") + str(nid),
        "标题": g("标题"), "链接": g("帖子链接"), "行业": g("行业"), "日期": day,
        "样本ID": g("样本ID"), "资格": g("资格"), "作者": g("作者"), "发布日期原样": g("发布日期原样"),
        "互动数据": likes.group(1) if likes else "待补", "拆解图片数": g("拆解图片数"),
        "图序观察": g("图序观察"), "设计结论": g("设计结论"), "点击机制_推断": g("点击机制_推断"),
        "咨询解释": g("咨询解释"), "本品牌迁移": g(f"{brand}迁移", "本品牌迁移"), "不可照搬": g("不可照搬"),
        "查看边界": g("查看边界"), "标签": g("标签"),
        "可借鉴点": section(md, f"可借鉴到{brand}的点") or section(md, "可借鉴点"),
        "图片文件": "；".join(str(folder / i) for i in imgs) or "待补",
        "图片URL": "；".join(urls) or "待补",
    }
    for k, v in (cfg.get("defaults") or {}).items(): row.setdefault(k, v)
    for f in S["core_fields"]:
        if "default" in f: row.setdefault(f["name"], f["default"])
    row["_md"] = md
    return row

def map_row(row, cfg, S, target_cols):
    """标准字段 → 目标列。映射不到的拆解字段并进 内容。"""
    fm = cfg.get("field_map") or {}
    out, extra = {}, []
    for k in all_fields(S):
        v = row.get(k)
        if k == S["content_field"] or v in (None, ""): continue
        col = fm.get(k, k) if k in fm else k
        if col and col in target_cols: out[col] = str(v)
        elif k in S["teardown_fields"]: extra.append(k)
    cc = fm.get(S["content_field"], S["content_field"])
    body = row["_md"]
    tail = [f"图片文件：{row['图片文件']}", f"图片URL：{row['图片URL']}",
            "（目标库无附件列：图片未上传，只记录本机路径和原始URL；原图仅作设计研究，不作商业素材）"]
    body += "\n\n## 图片\n" + "\n".join(tail)
    if cc in target_cols: out[cc] = body[:90000]
    missing = [c for c in out if c not in target_cols]
    return out, missing

# ---------- 后端 ----------
class CSVKB:
    def __init__(s, path, S, xlsx=False):
        s.p = Path(os.path.expanduser(path)); s.S = S; s.xlsx = xlsx
    def exists(s): return s.p.exists()
    def init(s):
        if s.p.exists(): sys.exit(f"已存在：{s.p}（不覆盖）")
        s.p.parent.mkdir(parents=True, exist_ok=True)
        cols = all_fields(s.S)
        if s.xlsx:
            import openpyxl; wb = openpyxl.Workbook(); ws = wb.active; ws.title = "对标拆解"; ws.append(cols); wb.save(s.p)
        else:
            with open(s.p, "w", newline="", encoding="utf-8-sig") as f: csv.writer(f).writerow(cols)
        return str(s.p)
    def cols(s):
        if s.xlsx:
            import openpyxl; ws = openpyxl.load_workbook(s.p).active
            return [c.value for c in ws[1]]
        with open(s.p, encoding="utf-8-sig") as f: return next(csv.reader(f))
    def ids(s):
        idf = (s.S["id_field"])
        if s.xlsx:
            import openpyxl; ws = openpyxl.load_workbook(s.p).active; cols = [c.value for c in ws[1]]
            i = cols.index(idf); return {r[i] for r in ws.iter_rows(min_row=2, values_only=True)}
        with open(s.p, encoding="utf-8-sig") as f: return {r.get(idf) for r in csv.DictReader(f)}
    def add(s, rec):
        cols = s.cols()
        if s.xlsx:
            import openpyxl; wb = openpyxl.load_workbook(s.p); wb.active.append([rec.get(c, "") for c in cols]); wb.save(s.p)
        else:
            with open(s.p, "a", newline="", encoding="utf-8-sig") as f: csv.writer(f).writerow([rec.get(c, "") for c in cols])
        return str(s.p)
    def read(s, rid):
        if s.xlsx:
            import openpyxl; ws = openpyxl.load_workbook(s.p).active; cols = [c.value for c in ws[1]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                d = dict(zip(cols, r))
                if d.get(s.S["id_field"]) == rid: return d
        else:
            with open(s.p, encoding="utf-8-sig") as f:
                for d in csv.DictReader(f):
                    if d.get(s.S["id_field"]) == rid: return d

class FeishuKB:
    B = "https://open.feishu.cn/open-apis"
    def __init__(s, fc, S):
        s.fc = fc; s.S = S; s.app = fc.get("app_token"); s.tbl = fc.get("table_id"); s._t = None; s.h = None
        if fc.get("helper"):
            sp = importlib.util.spec_from_file_location("kb_fs_helper", fc["helper"])
            s.h = importlib.util.module_from_spec(sp); sp.loader.exec_module(s.h)
    def _tok(s):
        if not s._t:
            body = json.dumps({"app_id": os.environ["FEISHU_APP_ID"], "app_secret": os.environ["FEISHU_APP_SECRET"]}).encode()
            r = json.load(urllib.request.urlopen(urllib.request.Request(s.B + "/auth/v3/tenant_access_token/internal", body, {"Content-Type": "application/json"})))
            assert r.get("code") == 0, r.get("msg"); s._t = r["tenant_access_token"]
        return s._t
    def api(s, m, path, params=None, body=None):
        if s.h: return s.h.api(m, path, params=params, json=body) if body is not None else s.h.api(m, path, params=params)
        url = s.B + path + ("?" + urllib.parse.urlencode(params) if params else "")
        req = urllib.request.Request(url, json.dumps(body).encode() if body is not None else None,
                                     {"Authorization": "Bearer " + s._tok(), "Content-Type": "application/json; charset=utf-8"}, method=m)
        try: return json.load(urllib.request.urlopen(req, timeout=60))
        except urllib.error.HTTPError as e: return json.loads(e.read() or b"{}")
    def init(s, name):
        if s.app: sys.exit("配置里已有 app_token，不再新建（避免重复建库）")
        j = s.api("POST", "/bitable/v1/apps", body={"name": name}); assert j.get("code") == 0, j
        app = j["data"]["app"]["app_token"]
        fields = []
        for f in s.S["core_fields"]:
            if f.get("type") == "select": fields.append({"field_name": f["name"], "type": 3, "property": {"options": [{"name": o} for o in f["options"]]}})
            else: fields.append({"field_name": f["name"], "type": 1})
        fields += [{"field_name": n, "type": 1} for n in s.S["teardown_fields"]]
        j = s.api("POST", f"/bitable/v1/apps/{app}/tables", body={"table": {"name": "对标与热点拆解", "default_view_name": "全部", "fields": fields}})
        assert j.get("code") == 0, j
        return {"app_token": app, "table_id": j["data"]["table_id"], "url": j["data"].get("url") or ""}
    def cols(s):
        j = s.api("GET", f"/bitable/v1/apps/{s.app}/tables/{s.tbl}/fields", params={"page_size": 100}); assert j.get("code") == 0, j
        return [f["field_name"] for f in j["data"]["items"]]
    def ids(s):
        out, pt = set(), None
        while True:
            p = {"page_size": 500, "field_names": json.dumps([s.S["id_field"]])}
            if pt: p["page_token"] = pt
            j = s.api("GET", f"/bitable/v1/apps/{s.app}/tables/{s.tbl}/records", params=p); assert j.get("code") == 0, j
            for it in j["data"].get("items") or []:
                v = it["fields"].get(s.S["id_field"])
                out.add(v if isinstance(v, str) else "".join(x.get("text", "") for x in v) if isinstance(v, list) else v)
            if not j["data"].get("has_more"): return out
            pt = j["data"]["page_token"]
    def add(s, rec):
        j = s.api("POST", f"/bitable/v1/apps/{s.app}/tables/{s.tbl}/records", body={"fields": rec})
        assert j.get("code") == 0, (j.get("code"), j.get("msg"))
        return j["data"]["record"]["record_id"]
    def read(s, rid_rec):
        j = s.api("GET", f"/bitable/v1/apps/{s.app}/tables/{s.tbl}/records/{rid_rec}", params={"with_shared_url": "true"})
        assert j.get("code") == 0, j; return j["data"]["record"]
    def link(s, rec_id, shared=None):
        return shared or (f"{s.fc.get('base_url')}?table={s.tbl}&record={rec_id}" if s.fc.get("base_url") else rec_id)

def backend(cfg, S):
    b = cfg.get("backend", "csv")
    if b == "feishu": return FeishuKB(cfg.get("feishu") or {}, S)
    if b == "xlsx": return CSVKB(cfg.get("xlsx_path", "~/xhs-tools/对标库/对标拆解库.xlsx"), S, xlsx=True)
    return CSVKB(cfg.get("csv_path", "~/xhs-tools/对标库/对标拆解库.csv"), S)

def main():
    import urllib.parse  # noqa
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["init", "check", "write"])
    ap.add_argument("folders", nargs="*")
    ap.add_argument("--config"); ap.add_argument("--backend"); ap.add_argument("--name", default="对标拆解库")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    cfg = load_cfg(a.config); S = schema()
    if a.backend: cfg["backend"] = a.backend
    kb = backend(cfg, S)
    if a.cmd == "init":
        if isinstance(kb, FeishuKB):
            r = kb.init(a.name); print("已新建飞书多维表格：", json.dumps(r, ensure_ascii=False))
            print("把 app_token / table_id / base_url 填进 config/kb_local.yaml，backend: feishu")
        else: print("已新建：", kb.init())
        return
    if isinstance(kb, CSVKB) and not kb.exists():
        sys.exit(f"还没有对标库：{kb.p}\n先问用户有没有自己的库；没有就运行 python scripts/kb.py init 按标准结构新建。")
    if isinstance(kb, FeishuKB) and not (kb.app and kb.tbl):
        sys.exit("飞书配置缺 app_token / table_id，见 README「接飞书多维表格」")
    cols = kb.cols()
    print(f"目标库（{cfg.get('backend','csv')}）列：", "、".join(cols))
    fm = cfg.get("field_map") or {}
    print("映射（标准字段 → 目标列；— 表示并进内容）：")
    for k in all_fields(S):
        col = fm.get(k, k) if k in fm else k
        print(f"  {k} → {col if col and col in cols else '—'}")
    if S["id_field"] not in cols: sys.exit(f"目标库没有去重列「{S['id_field']}」，请在 field_map 映射或让用户确认后加列（本脚本不改表结构）")
    if a.cmd == "check": return
    if not cfg.get("confirmed_mapping"): sys.exit("映射还没经用户确认（confirmed_mapping: false）。先把上面的映射发给用户确认，再改成 true。")
    have = kb.ids(); rc = 0
    for fo in a.folders:
        try: row = build_row(fo, cfg, S)
        except ValueError as e: print("✗", e); rc = 1; continue
        rid = row[S["id_field"]]
        if rid in have: print(f"↷ 跳过：{rid} 已在库里（不覆盖、不修改）"); continue
        rec, _ = map_row(row, cfg, S, cols)
        if a.dry_run: print("（预演）将新增：", rid, list(rec)); continue
        got = kb.add(rec); have.add(rid)
        if isinstance(kb, FeishuKB):
            back = kb.read(got); f = back["fields"]
            ok = all((f.get(k) if isinstance(f.get(k), str) else json.dumps(f.get(k), ensure_ascii=False)) and str(rec[k])[:50] in (f.get(k) if isinstance(f.get(k), str) else json.dumps(f.get(k), ensure_ascii=False)) for k in (S["id_field"], "标题"))
            print(f"✓ 已新增 {rid}（record_id={got}）读回校验：{'通过' if ok else '不一致'}")
            print("  链接：", kb.link(got, back.get("shared_url")))
        else:
            ok = (kb.read(rid) or {}).get("标题") == rec.get("标题")
            print(f"✓ 已新增 {rid} → {got} 读回校验：{'通过' if ok else '不一致'}")
    sys.exit(rc)

if __name__ == "__main__":
    import urllib.parse
    main()
